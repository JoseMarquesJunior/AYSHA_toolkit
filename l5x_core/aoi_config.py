"""Export of the configuration stored in AOI instance tags.

Every instance tag of an Add-On Instruction carries, in its decorated data, the
value of each parameter and local tag (ranges, setpoints, delays, HMI flags...).
Studio 5000 offers no practical way to list them for all instances at once; this
module builds one table per AOI type (one row per instance, one column per
member) and writes it as an Excel workbook (one sheet per AOI, plus a summary
and a data dictionary) or CSV.

Modes:
  "config": Input parameters that are not operands of the call (they live in the
            instance tag) + local tags editable by the HMI (ExternalAccess
            Read/Write); timers/counters keep only PRE.
  "all":    every parameter (except EnableIn/EnableOut) and every local tag,
            fully flattened.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Tuple

from .model import Aoi, Controller, Tag
from .xref import CrossRef, build_xref

MODES = ("config", "all")
FIXED_COLUMNS = ["PLC", "AOI", "Instancia", "Escopo", "Descricao", "Unidade", "Area"]
_TIMER_LIKE = {"TIMER", "COUNTER"}


# --------------------------------------------------------------------------
# Flattening
# --------------------------------------------------------------------------
def flatten(values: Any, prefix: str = "") -> Dict[str, Any]:
    """Nested decorated values -> {'Wrk_T.PRE': 5000, 'Arr[0].X': 1, ...}.

    STRING structures ({LEN, DATA}) collapse to their text."""
    out: Dict[str, Any] = {}
    if isinstance(values, dict):
        if set(values) == {"LEN", "DATA"}:
            out[prefix] = values.get("DATA")
            return out
        for key, val in values.items():
            path = f"{prefix}{key}" if str(key).startswith("[") else (f"{prefix}.{key}" if prefix else str(key))
            out.update(flatten(val, path))
        return out
    if prefix:
        out[prefix] = values
    return out


def _is_array_of_instances(tag: Tag) -> bool:
    return bool(tag.dimensions) and isinstance(tag.values, dict) and all(str(k).startswith("[") for k in tag.values)


# --------------------------------------------------------------------------
# Column selection
# --------------------------------------------------------------------------
@dataclass
class MemberInfo:
    name: str
    kind: str  # "parâmetro" | "local"
    usage: str  # Input | Output | InOut | local
    data_type: str
    required: bool = False
    visible: bool = False
    external_access: str = ""
    default: Any = None
    description: str = ""


def members_for(aoi: Aoi, mode: str) -> List[MemberInfo]:
    """Members of the AOI definition exported in ``mode``, in definition order."""
    out: List[MemberInfo] = []
    for p in aoi.parameters:
        if p.is_system:
            continue
        if mode == "config" and not (p.usage == "Input" and not p.required):
            continue
        out.append(MemberInfo(p.name, "parâmetro", p.usage, p.data_type, p.required, p.visible, p.external_access or "", p.default, p.description or ""))
    for t in aoi.local_tags:
        if t.usage:
            continue  # parameters were listed above
        if mode == "config" and not t.hmi_writable:
            continue
        out.append(MemberInfo(t.name, "local", "local", t.data_type + (f"[{t.dimensions}]" if t.dimensions else ""), False, False, t.external_access or "", t.values, t.description or ""))
    return out


def _member_columns(member: MemberInfo, sample: Dict[str, Any], mode: str) -> List[str]:
    """Flattened column names of one member, discovered from the instance values."""
    base = member.name
    cols = [c for c in sample if c == base or c.startswith(base + ".") or c.startswith(base + "[")]
    if mode == "config" and member.data_type.upper() in _TIMER_LIKE:
        cols = [c for c in cols if c.endswith(".PRE")]
    return cols or [base]


# --------------------------------------------------------------------------
# Tables
# --------------------------------------------------------------------------
@dataclass
class ConfigSheet:
    aoi: str
    columns: List[str]  # member columns (after the fixed ones)
    rows: List[Dict[str, Any]] = field(default_factory=list)
    dictionary: List[MemberInfo] = field(default_factory=list)
    protected: bool = False

    @property
    def header(self) -> List[str]:
        return FIXED_COLUMNS + self.columns


def _instances(ctrl: Controller, aoi: Aoi, xref: CrossRef) -> List[Tuple[Tag, str, Dict[str, Any]]]:
    """(tag, instance name, flattened values) for each instance (array elements expanded)."""
    out: List[Tuple[Tag, str, Dict[str, Any]]] = []
    for tag in xref.instances_of(aoi.name):
        if tag.scope.startswith("aoi:"):
            continue  # instances nested inside other AOIs have no own data
        if tag.values is None:
            out.append((tag, tag.name, {}))
            continue
        if _is_array_of_instances(tag):
            for idx, val in tag.values.items():
                out.append((tag, f"{tag.name}{idx}", flatten(val)))
        else:
            out.append((tag, tag.name, flatten(tag.values)))
    return out


def build_sheets(ctrl: Controller, xref: Optional[CrossRef] = None, mode: str = "config", aoi: Optional[str] = None) -> List[ConfigSheet]:
    """One ConfigSheet per AOI type that has instances (or only ``aoi``)."""
    if mode not in MODES:
        raise ValueError(f"mode must be one of {MODES}")
    xref = xref or build_xref(ctrl)
    sheets: List[ConfigSheet] = []
    for a in ctrl.aois:
        if aoi and a.name.lower() != aoi.lower():
            continue
        instances = _instances(ctrl, a, xref)
        if not instances:
            continue
        sample: Dict[str, Any] = {}
        for _, _, vals in instances:
            for k in vals:
                sample.setdefault(k, None)
        if a.protected:
            # no definition: export whatever the instance data holds
            columns = list(sample)
            dictionary: List[MemberInfo] = []
        else:
            dictionary = members_for(a, mode)
            columns = []
            for m in dictionary:
                for c in _member_columns(m, sample, mode):
                    if c not in columns:
                        columns.append(c)
        sheet = ConfigSheet(aoi=a.name, columns=columns, dictionary=dictionary, protected=a.protected)
        for tag, name, vals in instances:
            row: Dict[str, Any] = {
                "PLC": ctrl.name, "AOI": a.name, "Instancia": name,
                "Escopo": "controlador" if tag.scope == "controller" else tag.scope,
                "Descricao": " ".join((tag.description or "").split()), "Unidade": tag.engineering_unit or "", "Area": tag.area or "",
            }
            for c in columns:
                row[c] = vals.get(c, "")
            sheet.rows.append(row)
        sheets.append(sheet)
    return sheets


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------
def to_csv(sheet: ConfigSheet) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow(sheet.header)
    for r in sheet.rows:
        w.writerow([_cell(r.get(c, "")) for c in sheet.header])
    return buf.getvalue()


def _cell(v: Any) -> Any:
    if isinstance(v, (dict, list)):
        return str(v)
    return v


def _sheet_title(name: str, used: set) -> str:
    title = "".join(ch if ch not in '[]:*?/\\' else "_" for ch in name)[:31] or "AOI"
    base, n = title, 2
    while title.lower() in used:
        suffix = f"~{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def write_xlsx(sheets: Iterable[ConfigSheet], target: Any, mode: str = "config") -> None:
    """Excel workbook: 'Resumo', one sheet per AOI, and 'Dicionario' (member metadata).
    ``target`` is a path or a binary file-like object."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    sheets = list(sheets)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["AOI", "Instancias", "Colunas", "Protegida", "Modo"])
    for s in sheets:
        ws.append([s.aoi, len(s.rows), len(s.columns), "sim" if s.protected else "", mode])
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    used: set = {"resumo", "dicionario"}
    for s in sheets:
        wss = wb.create_sheet(_sheet_title(s.aoi, used))
        wss.append(s.header)
        for cell in wss[1]:
            cell.font = bold
        for r in s.rows:
            wss.append([_cell(r.get(c, "")) for c in s.header])
        wss.freeze_panes = "D2"
        for i, col in enumerate(s.header, 1):
            width = max([len(str(col))] + [len(str(r.get(col, ""))) for r in s.rows[:200]])
            wss.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
    wd = wb.create_sheet("Dicionario")
    wd.append(["AOI", "Membro", "Tipo de membro", "Uso", "DataType", "Obrigatorio", "Visivel", "ExternalAccess", "Default", "Descricao"])
    for cell in wd[1]:
        cell.font = bold
    for s in sheets:
        for m in s.dictionary:
            wd.append([s.aoi, m.name, m.kind, m.usage, m.data_type, "sim" if m.required else "", "sim" if m.visible else "", m.external_access, _cell(m.default) if m.default is not None else "", m.description])
    wb.save(target)


def xlsx_bytes(sheets: Iterable[ConfigSheet], mode: str = "config") -> bytes:
    buf = io.BytesIO()
    write_xlsx(sheets, buf, mode)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Generic tables (used by the alarm list, I/O list, drift and diff exports)
# --------------------------------------------------------------------------
Table = Tuple[str, List[str], List[Dict[str, Any]]]  # (title, header, rows)


def table_csv(header: List[str], rows: Iterable[Dict[str, Any]]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow(header)
    for r in rows:
        w.writerow([_cell(r.get(c, "")) for c in header])
    return buf.getvalue()


def write_tables_xlsx(tables: Iterable[Table], target: Any) -> None:
    """One sheet per (title, header, rows); ``target`` is a path or binary file-like."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    bold = Font(bold=True)
    for title, header, rows in tables:
        ws = wb.create_sheet(_sheet_title(title, used))
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        rows = list(rows)
        for r in rows:
            ws.append([_cell(r.get(c, "")) for c in header])
        ws.freeze_panes = "A2"
        for i, col in enumerate(header, 1):
            width = max([len(str(col))] + [len(str(r.get(col, ""))) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
    if not wb.sheetnames:
        wb.create_sheet("vazio")
    wb.save(target)


def tables_xlsx_bytes(tables: Iterable[Table]) -> bytes:
    buf = io.BytesIO()
    write_tables_xlsx(tables, buf)
    return buf.getvalue()
