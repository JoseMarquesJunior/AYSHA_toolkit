"""Lista de I/O e instrumentos, lista de alarmes, drift e diff semântico (fixture plant.l5x)."""

import os
import re
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from l5x_core import parse_file, parse_bytes, build_xref, analyze
from l5x_core.alarms import alarm_markdown, alarm_rows, alarm_summary, is_alarm_aoi
from l5x_core.aoi_config import table_csv, write_tables_xlsx
from l5x_core.diff import diff_markdown, diff_projects, diff_rows
from l5x_core.drift import drift_markdown, instance_drift, library_drift, library_markdown
from l5x_core.io_list import channel_summary, instruments, io_channels, io_source_map, io_tables, range_kind
from app.llm import ProjectContext
from app.streamlit_app import UI
from conftest import ROOT

PLANT = os.path.join(ROOT, "tests", "data", "plant.l5x")


@pytest.fixture(scope="module")
def pctrl():
    return parse_file(PLANT)


@pytest.fixture(scope="module")
def pxref(pctrl):
    return build_xref(pctrl)


def _modified_plant() -> bytes:
    """Versão 'nova' do fixture: range alterado, rung alterado, rotina removida, tag nova, AOI em nova revisão."""
    s = open(PLANT, encoding="utf-8").read()
    s = s.replace('<Tag Name="LIT_100"', '<Tag Name="LIT_100"', 1)
    s = re.sub(r'(<Tag Name="LIT_100".*?Cfg_RangeMax" DataType="REAL" Radix="Float" Value=")100\.0', r"\g<1>150.0", s, count=1, flags=re.S)
    s = s.replace("XIC(Local:3:I.Data.1)OTE(ZSH_100);", "XIC(Local:3:I.Data.1)XIO(XV_100_Out)OTE(ZSH_100);")
    s = re.sub(r'<Routine Name="Spare" Type="RLL">.*?</Routine>\n', "", s, flags=re.S)
    s = s.replace('<Tag Name="Rec1"', '<Tag Name="NovaTag" TagType="Base" DataType="DINT" Radix="Decimal" Constant="false" ExternalAccess="Read/Write"/>\n<Tag Name="Rec1"')
    s = s.replace('<AddOnInstructionDefinition Name="P_Alm" Revision="1.0"', '<AddOnInstructionDefinition Name="P_Alm" Revision="1.1"')
    s = s.replace('Value="55"', 'Value="99"')  # ACC do timer: runtime, não deve aparecer no diff
    s = s.replace('ExportDate="Wed Oct 14 09:00:00 2026"', 'ExportDate="Thu Oct 15 09:00:00 2026"')
    return s.encode("utf-8")


# --------------------------------------------------------------------------
# I/O e instrumentos
# --------------------------------------------------------------------------
def test_module_channels_parsed(pctrl):
    di = next(m for m in pctrl.modules if m.name == "DI_S3")
    assert di.slot == "3" and di.io_base == "Local:3" and di.major == "3"
    assert [(c.direction, c.path, c.description, c.spare) for c in di.channels] == [
        ("I", "Local:3:I.DATA.0", "WIRED SPARE", True), ("I", "Local:3:I.DATA.1", "ZSH-100 fim de curso aberto", False), ("I", "Local:3:I.DATA.2", "ZSL-100 fim de curso fechado", False),
    ]
    ai = next(m for m in pctrl.modules if m.name == "AI_S4")
    ch0 = ai.channels[0]
    assert ch0.engineering_unit == "m" and ch0.min == 4 and ch0.max == 20
    assert [c.path for c in ai.channels] == ["Local:4:I.CH0.DATA", "Local:4:I.CH1.DATA", "Local:4:I.CH2.DATA", "Local:4:O.CH0.DATA"]  # CH2 só tem unidade
    assert next(m for m in pctrl.modules if m.name == "Local").io_base is None


def test_io_channels_and_usage(pctrl, pxref):
    rows = {r["Canal"]: r for r in io_channels(pctrl, pxref)}
    assert rows["Local:3:I.DATA.1"]["Usos"] == 1 and rows["Local:3:I.DATA.1"]["Rotinas"] == "Main/MainRoutine"
    assert rows["Local:4:I.CH1.DATA"]["Rotinas"] == "Main/Scaling"  # referência em FBD
    assert rows["Local:4:O.CH0.DATA"]["Direcao"] == "O" and rows["Local:4:O.CH0.DATA"]["Usos"] == 1
    assert rows["Local:3:I.DATA.2"]["Usos"] == 0 and rows["Local:3:I.DATA.0"]["Spare"] == "sim"
    assert channel_summary(list(rows.values())) == {"canais": 7, "spare": 1, "sem_uso": 2, "sem_descricao": 1, "modulos": 2}


def test_instruments_link_channels(pctrl, pxref):
    source, sink = io_source_map(pctrl)
    assert source["lit_100_raw"] == "Local:4:I.Ch0.Data"  # MOV
    assert source["lit_200_raw"] == "Local:4:I.Ch1.Data"  # fio FBD
    assert source["zsh_100"] == "Local:3:I.Data.1"  # XIC -> OTE no mesmo rung
    assert sink["xv_100_out"] == "Local:4:O.Ch0.Data"
    rows = {r["Tag"]: r for r in instruments(pctrl, pxref)}
    assert set(rows) == {"LIT_100", "LIT_200", "LIT_300", "LIT_400", "LIT_500", "XV_100"}  # XV_200 sem unidade/range/canal fica fora
    lit = rows["LIT_100"]
    assert (lit["Unidade"], lit["Area"], lit["RangeMin"], lit["RangeMax"]) == ("m", "Tanques", 0.0, 100.0)
    assert lit["EntradaIO"] == "Local:4:I.Ch0.Data" and lit["EntradaIODescricao"] == "LIT-100 nível tanque 1" and lit["Chamada"] == "Main/MainRoutine rung 0"
    assert rows["LIT_200"]["EntradaIO"] == "Local:4:I.Ch1.Data"
    assert rows["XV_100"]["EntradaIO"] == "Local:3:I.Data.1" and rows["XV_100"]["SaidaIO"] == "Local:4:O.Ch0.Data"
    assert "XV_200" in {r["Tag"] for r in instruments(pctrl, pxref, include_all=True)}
    assert range_kind("rRangeMin") == "min" and range_kind("EU_Max") == "max" and range_kind("bEnableUnderRangeAlarm") is None and range_kind("Raw_Min") is None
    tables = io_tables(pctrl, pxref)
    assert [t[0] for t in tables] == ["Instrumentos", "Canais_IO"]


# --------------------------------------------------------------------------
# Alarmes
# --------------------------------------------------------------------------
def test_alarm_list(pctrl, pxref):
    assert is_alarm_aoi(pctrl.aoi("P_Alm")) and not is_alarm_aoi(pctrl.aoi("AI_Std"))
    header, rows = alarm_rows(pctrl, pxref)
    assert header[:10] == ["Alarme", "Tipo", "Escopo", "Instancia", "InstanciaTipo", "Descricao", "DescricaoMembro", "Severidade", "Habilitado", "Uso"]
    assert {"MinDurationPRE", "AckRequired", "Latched", "Cfg_AckReqd"} <= set(header) and "InAlarm" not in header
    by = {r["Alarme"]: r for r in rows}
    assert set(by) == {"DigAlm1", "LIT_100.AH_Alm", "LIT_200.AH_Alm", "LIT_300.AH_Alm", "LIT_400.AH_Alm", "LIT_500.AH_Alm"}
    d = by["DigAlm1"]
    assert (d["Tipo"], d["Severidade"], d["AckRequired"], d["Uso"]) == ("ALMD", 600, 1, "Main/MainRoutine")
    a = by["LIT_100.AH_Alm"]
    assert (a["Tipo"], a["InstanciaTipo"], a["Severidade"], a["Habilitado"], a["Cfg_AckReqd"]) == ("P_Alm", "AI_Std", 750, 1, 1)
    assert a["Descricao"] == "Nível tanque 1" and a["DescricaoMembro"] == "Alarme de nível alto" and a["Uso"] == "Main/MainRoutine rung 0"
    assert by["LIT_200.AH_Alm"]["Habilitado"] == 0 and by["LIT_300.AH_Alm"]["Uso"] == ""
    s = alarm_summary(rows)
    assert s == {"total": 6, "por_tipo": {"ALMD": 1, "P_Alm": 5}, "por_severidade": {"600": 1, "750": 1, "500": 4}, "desabilitados": 1, "sem_descricao": 0, "sem_uso": 3, "descricoes_duplicadas": 0}
    md = alarm_markdown(header, rows)
    assert md.startswith("# Lista de alarmes") and "| LIT_100.AH_Alm | P_Alm |" in md and "Desabilitados (Cfg_Exists=0): 1" in md


# --------------------------------------------------------------------------
# Drift
# --------------------------------------------------------------------------
def test_instance_drift(pctrl, pxref):
    rows = instance_drift(pctrl, pxref)
    assert {(r["AOI"], r["Membro"], r["ValorComum"], r["Detalhe"]) for r in rows} == {
        ("AI_Std", "Cfg_RangeMax", "100.0", "LIT_400=250.0"),
        ("AI_Std", "AH_Alm.Cfg_Severity", "500", "LIT_100=750"),
        ("AI_Std", "AH_Alm.Cfg_Exists", "1", "LIT_200=0"),
        ("AI_Std", "AH_Alm.Cfg_AckReqd", "1", "LIT_500=0"),
    }
    assert all(r["Participacao"] == "80%" and r["Instancias"] == 5 for r in rows)
    assert instance_drift(pctrl, pxref, min_instances=6) == []  # XV_Std só tem 2 instâncias, AI_Std 5
    assert "LIT_400=250.0" in drift_markdown(rows)


def test_library_drift(pctrl):
    new = parse_bytes(_modified_plant())
    rows = {(r["Tipo"], r["Item"]): r for r in library_drift([pctrl, new])}
    assert rows[("AOI", "P_Alm")]["Status"] == "revisão diferente" and "PLANT: rev. 1.0" in rows[("AOI", "P_Alm")]["Detalhe"]
    assert rows[("AOI", "AI_Std")]["Status"] == "igual" and rows[("UDT", "Recipe")]["Status"] == "igual"
    # interface diferente e ausência
    other = parse_bytes(open(PLANT, encoding="utf-8").read().replace('<Parameter Name="Cfg_TravelT"', '<Parameter Name="Cfg_TravelTime"').replace('<AddOnInstructionDefinition Name="AI_Std"', '<AddOnInstructionDefinition Name="AI_Std2"').encode("utf-8"))
    rows = {(r["Tipo"], r["Item"]): r for r in library_drift([pctrl, other])}
    assert rows[("AOI", "XV_Std")]["Status"] == "interface diferente" and "cfg_travelt" in rows[("AOI", "XV_Std")]["Detalhe"]
    assert rows[("AOI", "AI_Std")]["Status"] == "ausente em PLANT" or rows[("AOI", "AI_Std")]["Status"].startswith("ausente em")
    md = library_markdown(list(rows.values()), ["PLANT", "PLANT"])
    assert "| XV_Std | AOI | interface diferente |" in md


# --------------------------------------------------------------------------
# Diff
# --------------------------------------------------------------------------
def test_diff_projects(pctrl, pxref):
    new = parse_bytes(_modified_plant())
    old_f, _ = analyze(pctrl, pxref)
    new_f, _ = analyze(new, build_xref(new))
    d = diff_projects(pctrl, new, old_f, new_f)
    assert not d.empty and d.old_date.startswith("Wed") and d.new_date.startswith("Thu")
    by = d.by_category()
    assert [(c.kind, c.item) for c in by["rotina"]] == [("removido", "Main/Spare")]
    rung = by["rung"]
    assert len(rung) == 1 and rung[0].kind == "alterado" and rung[0].detail == "rung 2" and "XIO(XV_100_Out)" in rung[0].new
    assert [(c.kind, c.item) for c in by["tag"]] == [("adicionado", "controlador/NovaTag")]
    cfg = by["config"]
    assert [(c.item, c.detail, c.old, c.new) for c in cfg] == [("controlador/LIT_100", "Cfg_RangeMax", 100.0, 150.0)]  # ACC do timer ignorado
    assert [(c.kind, c.item, c.detail, c.old, c.new) for c in by["aoi"]] == [("alterado", "P_Alm", "revisão", "1.0", "1.1")]
    assert not by["programa"] and not by["udt"] and not by["modulo"] and not by["task"]
    findings = {(c.kind, c.detail) for c in by["finding"]}
    assert ("removido", "[warning] routine_never_called") in findings  # Spare sumiu
    md = diff_markdown(d)
    assert md.startswith("# Diff: PLANT (Wed Oct 14 09:00:00 2026) → PLANT (Thu Oct 15 09:00:00 2026)")
    assert "## Configuração armazenada nas tags (1)" in md and "| alterado | controlador/LIT_100 | Cfg_RangeMax | 100.0 | 150.0 |" in md
    rows = diff_rows(d)
    assert rows[0].keys() == {"Categoria", "Tipo", "Item", "Detalhe", "Antes", "Depois"}
    # mesmo arquivo: nenhuma diferença
    same = diff_projects(pctrl, parse_file(PLANT), old_f, old_f)
    assert same.empty and "Nenhuma diferença" in diff_markdown(same)


# --------------------------------------------------------------------------
# CLI, exports e app
# --------------------------------------------------------------------------
def test_cli_reports(tmp_path):
    new_path = tmp_path / "plant_v2.l5x"
    new_path.write_bytes(_modified_plant())
    run = lambda *a: subprocess.run([sys.executable, "-X", "utf8", "-m", "l5x_core.cli", PLANT, *a], capture_output=True, text=True, encoding="utf-8")  # noqa: E731
    p = run("--alarms", "--export-alarms", str(tmp_path / "alm.xlsx"))
    assert p.returncode == 0, p.stderr
    assert "# Lista de alarmes" in p.stdout and "6 linhas" in p.stderr and load_workbook(tmp_path / "alm.xlsx").sheetnames == ["Alarmes"]
    p = run("--io-list", "--export-io", str(tmp_path / "io.xlsx"))
    assert p.returncode == 0 and "6 instrumento(s); 7 canal(is)" in p.stdout and "Local:3:I.DATA.0" in p.stdout and "[SPARE]" in p.stdout
    assert load_workbook(tmp_path / "io.xlsx").sheetnames == ["Instrumentos", "Canais_IO"]
    p = run("--drift")
    assert p.returncode == 0 and "LIT_400=250.0" in p.stdout
    p = run("--library-drift", str(new_path))
    assert p.returncode == 0 and "| P_Alm | AOI | revisão diferente |" in p.stdout
    p = run("--diff", str(new_path), "--export-diff", str(tmp_path / "diff.csv"))
    assert p.returncode == 0 and "diferenças" in p.stderr
    csv = (tmp_path / "diff.csv").read_text(encoding="utf-8-sig")
    assert csv.splitlines()[0] == "Categoria;Tipo;Item;Detalhe;Antes;Depois" and "config;alterado;controlador/LIT_100;Cfg_RangeMax;100.0;150.0" in csv
    p = run("--diff", str(new_path))
    assert "# Diff: PLANT" in p.stdout
    p = run("--diff", str(tmp_path / "missing.l5x"))
    assert p.returncode == 2
    p = run("--export-alarms", str(tmp_path / "alm.csv"))
    assert p.returncode == 0 and (tmp_path / "alm.csv").read_text(encoding="utf-8-sig").startswith("Alarme;Tipo;")


def test_project_context_and_ui(pctrl):
    ctx = ProjectContext(pctrl)
    header, rows = ctx.alarm_table()
    assert len(rows) == 6 and ctx.alarm_table()[1] is rows  # cache
    assert [t[0] for t in ctx.io_tables()] == ["Instrumentos", "Canais_IO"]
    assert len(ctx.drift_rows()) == 4
    other = ProjectContext(parse_bytes(_modified_plant()))
    d = ctx.diff_to(other)
    assert d.counts()["config"] == 1
    assert any(r["Status"] == "revisão diferente" for r in ctx.library_drift_with([other]))
    assert table_csv(["A", "B"], [{"A": 1, "B": "x;y"}]).splitlines()[1] == '1;"x;y"'  # delimitador dentro do campo é escapado
    for lang in UI:
        assert "nav_compare" in UI[lang] and "alarms_title" in UI[lang] and "io_title" in UI[lang] and "drift_title" in UI[lang]
